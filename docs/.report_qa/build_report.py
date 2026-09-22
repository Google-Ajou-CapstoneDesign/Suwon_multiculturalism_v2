from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = Path(__file__).resolve().parents[1] / 'Local_Bridge_2026-09_구현내역.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(.8)
sec.header_distance = sec.footer_distance = Inches(.492)
# compact_reference_guide + memo_masthead.
# Named overrides: Korean font Malgun Gothic; compact monthly-report headings
# 11.5pt/8pt before/4pt after; title 20pt. Two-page Korean summary overrides:
# 0.8in margins, 10.5pt body, 1.10 lines, 4pt paragraph spacing, no grid snap.
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Header', 'Footer']:
    style = doc.styles[name]
    style.font.name = '맑은 고딕'
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '맑은 고딕')
    style.font.size = Pt(10.5)
    style.font.color.rgb = RGBColor.from_string('243247')
    fmt = style.paragraph_format
    fmt.space_before, fmt.space_after = Pt(0), Pt(4)
    fmt.line_spacing = 1.10
    snap = OxmlElement('w:snapToGrid'); snap.set(qn('w:val'), '0')
    style.element.get_or_add_pPr().append(snap)
    fmt.widow_control = True
for name, size in [('Heading 1', 15), ('Heading 2', 11.5), ('Heading 3', 11.5)]:
    style = doc.styles[name]
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string('2E74B5')
    style.paragraph_format.space_before = Pt(8)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.keep_with_next = True
doc.styles['Title'].font.size = Pt(20)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(6)
doc.styles['Subtitle'].font.size = Pt(9)
doc.styles['Subtitle'].font.italic = False
doc.styles['Subtitle'].font.color.rgb = RGBColor.from_string('627084')
for name in ['Header', 'Footer']:
    doc.styles[name].font.size = Pt(8)
    doc.styles[name].font.color.rgb = RGBColor.from_string('627084')
    doc.styles[name].paragraph_format.space_after = Pt(0)

header = sec.header.paragraphs[0]
header.text = 'LOCAL BRIDGE  |  월간 개발 현황'
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('2026.09.20 기준  ·  ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
footer.add_run(' / ')
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'NUMPAGES'); footer._p.append(field)

def section(title, text):
    doc.add_paragraph(title, 'Heading 2')
    doc.add_paragraph(text)

doc.add_paragraph('2026년 9월 구현 내역', 'Title')
doc.add_paragraph('수원시 이주민 노동·생활 정보 플랫폼  |  집계 기간: 9월 1일~20일', 'Subtitle')
doc.add_paragraph('9월에는 근무기록과 임금 확인, 증빙 업로드, AI 상담 및 다국어 이용 흐름을 확장했다. 아래는 9월 커밋 11건과 최신 작업본을 함께 확인한 기능별 요약이다.')

section('09.06  위치 인증에 주소·증빙 정보 추가',
    'GPS 좌표를 주소로 변환해 홈과 근무기록장에 표시하도록 개선했다. 인증 당시 위도·경도·주소를 근무기록에 함께 저장하며, 주소 변환이 실패해도 위치 인증은 유지한다. 주소 조회에는 사용자가 선택한 언어를 전달한다.')
section('09.10  화면 구조·온보딩·설정 기능 정비',
    '홈·임금계산기·네비게이터·설정 중심으로 화면을 재구성하고, 임금체불·산재 대응 진입 화면을 추가했다. 온보딩을 언어 선택으로 간소화하고 캘린더에 로그인/데모 선택 화면을 연결했다. 기능 위치를 짚는 스포트라이트 투어와 사용설명서, 프로필·국적·비자 편집, 로그아웃, 증빙보관함 진입을 구현했다. 알림 토글은 로컬 선호 설정이다. 테마·글꼴·레이아웃도 정비했다.')
section('09.10~14  근무기록·간편 임금 계산 개선',
    '캘린더에 일별 예상 임금과 월 합계를 추가하고, 간편 입력·오늘 예상 임금 카드·월 합계 배너로 구성했다. 시급과 근무시간을 조절하면 금액이 즉시 계산되며 기록을 저장할 수 있다. 홈의 이번 달 근무 카드에는 총근무시간 옆에 이번 달 임금을 표시하고, 로그인 시 서버 캘린더 기록으로 집계한다. 네비게이터의 불필요한 설명 문구도 정리했다.')
section('09.14  사업주 증빙보관함 업로드 연결',
    '근로계약서·임금명세서·사업주 메시지·통화 녹음을 파일 선택 후 업로드하고 목록에서 확인하도록 기존 API에 연결했다. 로그인 토큰을 전달하고 증빙 종류 및 근무 날짜로 구분한다. 실제 파일은 Storage, 사용자·종류·경로·업로드 시각 등의 메타데이터는 Firestore evidence_files에 보관한다.')
section('09.14  버그 리포트 기능 추가',
    '설정에서 버그 내용을 작성하고 teameqlab@gmail.com 앞으로 제목·본문이 채워진 메일 앱을 열도록 구현했다. 빈 입력을 검사하고 메일 앱 실행 실패 시 내용 복사를 제공한다. 서버 자동 발송이 아니라 사용자가 메일 앱에서 전송하는 방식이다.')

doc.add_page_break()
doc.add_paragraph('AI 가이드 고도화 및 최근 추가 사항', 'Heading 1')
section('09.15  AI 전송 제한·입력 보호',
    '비회원(IP)은 5초 간격·분당 5회·하루 10회, 로그인(UID)은 3초 간격·분당 10회·하루 50회로 제한했다. 동시 요청 1개·동일 IP 분당 60회, 한국 시간 자정 초기화와 Firestore 카운터 공유를 적용했다. 메시지 2,000자·이력 6개·합계 12,000자, AI 처리 60초 및 에이전트 호출 횟수를 제한하며, 제한 시 입력 보존·대기시간 안내를 제공한다.')
section('09.15  의도 분류를 한국 생활·기관 찾기로 확장',
    '임금·산재·근로계약·서비스 안내·범위 밖 요청에 한국 생활 정보(life_info)와 기관 찾기(org_search)를 더해 7개로 분류한다. 비자·의료·주거·교통·은행·교육 문의와 기관 위치·연락처·이용시간 질문을 처리하며, 이전 대화의 후속 질문도 참고한다. 새 분류를 에이전트 및 문서·기관 검색 도구에 연결했다.')
section('09.15~17  우즈베크어·터키어 지원',
    '우즈베크어(UZB)·터키어(TUR)를 추가해 총 6개 언어를 지원한다. 언어 선택, 화면·백과사전·네비게이터·계산기, AI 답변·기본 안내, 위치 조회와 PDF에 반영했다. 터키어의 월·요일·근속기간 표현도 보완했다. 터키어 변경분은 9월 20일 기준 작업본에 반영됐으며 미커밋 상태다.')
section('09.16  관련도 상위 2개 기관 상세 카드',
    '관련도순 Top-2 결과를 유지하도록 거리순 재정렬을 제거하고, 기관명·주소·전화번호·이용가능시간·거리를 API와 화면에 전달했다. 두 카드는 간격을 제외한 너비를 절반씩 사용한다. 긴 내용은 줄바꿈하고 누락된 정보는 별도 안내해 가로 스크롤로 카드가 잘리는 문제를 개선했다.')
section('09.16~17  업로드 설정·네비게이터 링크 정리',
    'Storage 설정 누락·버킷 없음·권한 오류를 구분해 503으로 안내한다. 배포 설정에 local-bridge-evidence 버킷과 환경변수 보존 방식을 반영했다. README에는 버킷 생성·권한 연결 완료가 기록돼 있다. 네비게이터의 동작하지 않는 백과사전 이동 버튼도 제거했다.')
section('09.17  챗봇 추천 질문 카드 추가',
    '빈 대화 화면에 임금·산재·지원기관 관련 추천 질문 3개를 표시하고, 선택하면 해당 언어의 질문을 바로 전송하도록 했다. 전송 제한 중에는 추천 질문도 비활성화한다.')

p = doc.add_paragraph('근거: 9월 Git 기록·README·현재 작업본(9/20). 날짜는 커밋일 중심이며 터키어는 9/17 작업 기록을 반영했다. 운영 배포·실서비스 동작은 이번 문서 작성에서 재검증하지 않았다.', 'Subtitle')
p.paragraph_format.space_before = Pt(6)
doc.core_properties.title = 'Local Bridge 2026년 9월 구현 내역'
doc.core_properties.author = 'Team EQ Lab'
doc.core_properties.subject = '2026-09-20 기준 월간 구현 현황'
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
